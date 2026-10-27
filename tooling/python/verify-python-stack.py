#!/usr/bin/env python3
"""Guncat Work: 原生 Python 技术栈自检（在**目标** python 里跑，通常经 qemu）。

用法（见 run-verify-stack.sh，一般不用手跑）：
    python3 verify-python-stack.py

逐项检查标准库、C 扩展、三方库，打印 PASS/FAIL 表；全部通过时返回 0。
每个检查都在 try 里跑，单项失败不影响其它项 —— 便于一次性看清"哪些能用、哪些缺"。
"""

from __future__ import annotations

import io
import os
import sys
import traceback

RESULTS = []


def check(name, fn):
    try:
        detail = fn()
        RESULTS.append((True, name, detail if detail is not None else ""))
    except BaseException as exc:  # noqa: BLE001 - 自检工具，任何异常都要记下来
        RESULTS.append((False, name, "%s: %s" % (type(exc).__name__, exc)))


# ---------- 标准库 / C 扩展 ----------
def _stdlib():
    import bz2, hashlib, json, lzma, sqlite3, ssl, zlib
    con = sqlite3.connect(":memory:")
    con.execute("create table t(x int)")
    con.execute("insert into t values (7)")
    total = con.execute("select sum(x) from t").fetchone()[0]
    payload = b"guncat" * 100
    assert zlib.decompress(zlib.compress(payload)) == payload
    assert bz2.decompress(bz2.compress(payload)) == payload
    assert lzma.decompress(lzma.compress(payload)) == payload
    assert hashlib.sha256(b"abc").hexdigest().startswith("ba7816bf")
    ctx = ssl.create_default_context()
    certs = ctx.get_ca_certs()
    # 应用启动时会把 SSL_CERT_FILE 指向随包的 cacert.pem; 有就该能读到根证书
    cafile = os.environ.get("SSL_CERT_FILE", "")
    if cafile:
        assert os.path.isfile(cafile), cafile
        assert len(certs) > 0, "SSL_CERT_FILE 已设置但没读到根证书"
        note = "cacert=%s(%d)" % (os.path.basename(cafile), len(certs))
    else:
        note = "CA=%d(未设置 SSL_CERT_FILE)" % len(certs)
    return "sqlite3=%s, %s, sha256 ok" % (sqlite3.sqlite_version, note)


def _ctypes():
    import ctypes
    libm = ctypes.CDLL("libm.so")
    libm.sqrt.restype = ctypes.c_double
    libm.sqrt.argtypes = [ctypes.c_double]
    return "sqrt(9)=%.1f" % libm.sqrt(9.0)


# ---------- 数值 / 数据 ----------
def _numpy():
    import numpy as np
    a = np.arange(12).reshape(3, 4)
    assert int(a.sum()) == 66
    fft = [round(float(x.real), 1) for x in np.fft.fft(np.arange(4))]
    assert fft == [6.0, -2.0, -2.0, -2.0], fft
    assert round(float(np.linalg.det(np.array([[1.0, 2.0], [3.0, 4.0]]))), 1) == -2.0
    # 必须真的启用 BLAS/LAPACK(真机上曾出现过"声明支持但 blas=none"的情况)
    cfg = np.show_config(mode="dicts")
    blas = cfg["Build Dependencies"]["blas"].get("name")
    lapack = cfg["Build Dependencies"]["lapack"].get("name")
    assert blas == "openblas" and lapack == "openblas", "numpy 未启用 BLAS/LAPACK: %s/%s" % (blas, lapack)
    return "%s blas=%s" % (np.__version__, blas)


def _scipy():
    """scipy 覆盖 linalg/optimize/special/stats/sparse/integrate 六个核心模块。"""
    import numpy as np
    from scipy import linalg, optimize, special, stats, sparse, integrate
    a = np.array([[3.0, 2.0], [1.0, 4.0]])
    x = linalg.solve(a, np.array([7.0, 9.0]))
    assert np.allclose(x, [1.0, 2.0]), x
    assert abs(float(linalg.det(a)) - 10.0) < 1e-9   # 3*4 - 2*1 = 10
    r = optimize.minimize(lambda v: (v[0] - 1) ** 2 + (v[1] - 2) ** 2, [0.0, 0.0])
    assert r.success and np.allclose(r.x, [1.0, 2.0], atol=1e-3), r.x
    assert abs(float(integrate.quad(lambda v: v * v, 0, 1)[0]) - 1.0 / 3.0) < 1e-6
    assert abs(float(special.erf(0.0))) < 1e-12
    assert abs(float(stats.norm.cdf(0.0)) - 0.5) < 1e-12
    assert sparse.csr_matrix(np.eye(3)).nnz == 3
    import scipy
    return scipy.__version__


def _pandas():
    import pandas as pd
    df = pd.DataFrame({"city": ["a", "b", "a", "b"], "amt": [1.5, 2.5, 3.0, 4.0]})
    got = df.groupby("city")["amt"].sum().to_dict()
    assert got == {"a": 4.5, "b": 6.5}, got
    ts = pd.date_range("2026-10-27", periods=2, freq="D", tz="Asia/Shanghai")
    assert "+08:00" in str(ts[0])
    return pd.__version__


# ---------- 图像 / 绘图 ----------
def _pillow():
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (32, 32), (10, 20, 30))
    ImageDraw.Draw(im).rectangle([4, 4, 20, 20], fill=(240, 220, 60))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    im2 = Image.open(io.BytesIO(buf.getvalue()))
    assert im2.size == (32, 32)
    import PIL
    return "%s png=%d bytes" % (PIL.__version__, len(buf.getvalue()))


def _matplotlib():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(2, 2))
    ax.plot([1, 2, 3], [2, 1, 3])
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    assert len(buf.getvalue()) > 1000
    return "%s png=%d bytes" % (matplotlib.__version__, len(buf.getvalue()))


def _cjk_font():
    """中文渲染: 拿不到字体时 matplotlib 会退回 DejaVu 并警告"Glyph missing", 这里必须为 0。"""
    import warnings
    import guncat_fonts
    font = guncat_fonts.path()
    assert font and os.path.isfile(font), "找不到随包中文字体 fonts/SourceHanSansSC-Regular.otf"
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        fig, ax = plt.subplots(figsize=(2, 1))
        ax.text(0.5, 0.5, "中文测试", ha="center", va="center", fontsize=16)
        ax.axis("off")
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        plt.close(fig)
    missing = [w for w in caught if "missing from font" in str(w.message)]
    assert not missing, "matplotlib 缺字: %s" % missing[0].message
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (120, 40), "white")
    ImageDraw.Draw(im).text((4, 8), "中文测试", font=guncat_fonts.pillow_font(20), fill="black")
    ink = sum(1 for v in im.convert("L").getdata() if v < 128)
    assert ink > 50, "Pillow 中文渲染没有墨迹(ink=%d)" % ink
    return "%s ink=%d" % (os.path.basename(font), ink)


def _qrcode():
    import qrcode
    img = qrcode.make("https://example.com")
    buf = io.BytesIO()
    img.save(buf)
    return "png=%d bytes" % len(buf.getvalue())


# ---------- 解析 / 文本 ----------
def _lxml():
    from lxml import etree, html
    doc = etree.fromstring('<root><a x="1">hi</a><a x="2">yo</a></root>')
    assert doc.xpath("//a/@x") == ["1", "2"]
    page = html.fromstring('<div><p class="c">A</p><p>B</p></div>')
    assert [e.text for e in page.cssselect("p.c")] == ["A"]
    import lxml
    return lxml.__version__


def _regex():
    import regex
    assert regex.findall(r"\p{Han}+", "中文abc汉字") == ["中文", "汉字"]
    assert regex.search(r"\p{Emoji}", "hi \U0001F30F")
    return regex.__version__


def _bs4():
    import bs4
    soup = bs4.BeautifulSoup("<p class='x'>hi</p>", "html.parser")
    assert soup.select_one("p.x").text == "hi"
    return bs4.__version__


# ---------- 加密 ----------
def _crypto():
    from Crypto.Cipher import AES
    from Crypto.Hash import SHA256
    from Crypto.PublicKey import RSA
    from Crypto.Signature import pkcs1_15
    key, nonce = b"0123456789abcdef0123456789abcdef", b"0123456789ab"
    cipher = AES.new(key, AES.MODE_GCM, nonce=nonce)
    ct, tag = cipher.encrypt_and_digest(b"hello")
    assert AES.new(key, AES.MODE_GCM, nonce=nonce).decrypt_and_verify(ct, tag) == b"hello"
    rsa = RSA.generate(2048)
    sig = pkcs1_15.new(rsa).sign(SHA256.new(b"m"))
    pkcs1_15.new(rsa.publickey()).verify(SHA256.new(b"m"), sig)
    import Crypto
    return "%s rsa=%d bits" % (Crypto.__version__, rsa.size_in_bits())


# ---------- 文档 ----------
def _docx():
    import docx
    d = docx.Document()
    d.add_heading("Guncat", 0)
    buf = io.BytesIO()
    d.save(buf)
    return "%d bytes" % len(buf.getvalue())


def _pptx():
    import pptx
    p = pptx.Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Slide"
    buf = io.BytesIO()
    p.save(buf)
    return "%d bytes" % len(buf.getvalue())


def _xlsxwriter():
    import xlsxwriter
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    wb.add_worksheet().write(0, 0, 42)
    wb.close()
    return "%d bytes" % len(buf.getvalue())


# ---------- 数学 / 系统 / 网络 ----------
def _sympy():
    import sympy
    x = sympy.symbols("x")
    assert str(sympy.integrate(x ** 2, x)) == "x**3/3"
    return sympy.__version__


def _psutil():
    import psutil
    assert psutil.cpu_count() >= 1
    return "%s cpu=%d threads=%d" % (psutil.__version__, psutil.cpu_count(),
                                     psutil.Process().num_threads())


def _requests():
    import requests
    import certifi
    assert os.path.isfile(certifi.where())
    return "%s certifi=%s" % (requests.__version__, certifi.__version__ if hasattr(certifi, "__version__") else "ok")


def _openpyxl_yaml_jinja():
    import openpyxl, yaml, jinja2
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    assert yaml.safe_load(yaml.safe_dump({"a": 1})) == {"a": 1}
    assert jinja2.Template("hi {{ n }}").render(n=1) == "hi 1"
    return "%s/%s/%s" % (openpyxl.__version__, yaml.__version__, jinja2.__version__)


def _reportlab():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("Helvetica", 12)
    c.drawString(72, 720, "Guncat")
    c.showPage()
    c.save()
    assert buf.getvalue().startswith(b"%PDF")
    return "pdf=%d bytes" % len(buf.getvalue())


def _watchdog():
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler
    obs = Observer()
    assert type(obs).__name__ in ("InotifyObserver", "Observer"), type(obs).__name__
    return "%s" % type(obs).__name__


CHECKS = [
    ("stdlib", _stdlib),
    ("ctypes", _ctypes),
    ("numpy", _numpy),
    ("pandas", _pandas),
    ("scipy", _scipy),
    ("pillow", _pillow),
    ("matplotlib", _matplotlib),
    ("cjk-font", _cjk_font),
    ("qrcode", _qrcode),
    ("lxml", _lxml),
    ("regex", _regex),
    ("bs4", _bs4),
    ("crypto", _crypto),
    ("docx", _docx),
    ("pptx", _pptx),
    ("xlsxwriter", _xlsxwriter),
    ("sympy", _sympy),
    ("psutil", _psutil),
    ("requests", _requests),
    ("openpyxl/yaml/jinja2", _openpyxl_yaml_jinja),
    ("reportlab", _reportlab),
    ("watchdog", _watchdog),
]


def main() -> int:
    print("python", sys.version.split()[0], "|", sys.platform)
    for name, fn in CHECKS:
        check(name, fn)
    passed = sum(1 for ok, _, _ in RESULTS if ok)
    print()
    for ok, name, detail in RESULTS:
        print(("  PASS  " if ok else "  FAIL  ") + "%-22s %s" % (name, detail))
    print("\n%d/%d passed" % (passed, len(RESULTS)))
    if passed != len(RESULTS):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
